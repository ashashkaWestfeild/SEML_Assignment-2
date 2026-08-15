"""Build a plain Word version of the report (``Group_84.docx``).

The content is taken from the same story that ``build_report.py`` renders to
PDF, not from the PDF's extracted text -- extracting from a PDF turns tables
into line soup and loses code indentation.

The output is deliberately unstyled: Word's built-in Normal/Heading styles, a
plain table grid, monospace for code, no colours, borders or custom fonts. It
is meant to be reformatted into whatever template is required.

Run:
    python scripts/build_word_report.py
"""

from __future__ import annotations

import html
import re
import sys
from pathlib import Path
from typing import Any, Iterable, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    Preformatted,
    Spacer,
    Table,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from build_report import build_story, load_artifacts  # noqa: E402

OUT = PROJECT_ROOT.parent / "Group_84.docx"

#: report style name -> Word built-in style
HEADING_STYLES = {
    "title": "Title",
    "h1": "Heading 1",
    "h2": "Heading 2",
    "h3": "Heading 3",
}
MAX_IMAGE_WIDTH_IN = 6.0


def to_plain_text(markup: str) -> str:
    """Strip the report's inline markup down to plain text."""
    text = re.sub(r"<br\s*/?>", " ", markup)
    text = re.sub(r"<[^>]+>", "", text)
    text = html.unescape(text)
    text = text.replace("≥", ">=").replace("≤", "<=")
    return re.sub(r"\s+", " ", text).strip()


def flatten(flowables: Iterable[Any]) -> List[Any]:
    """Expand KeepTogether wrappers into a flat sequence."""
    out: List[Any] = []
    for item in flowables:
        if isinstance(item, KeepTogether):
            out.extend(flatten(item._content))
        else:
            out.append(item)
    return out


def add_paragraph(document: Document, item: Paragraph) -> None:
    """Write one report paragraph using a built-in Word style."""
    style_name = getattr(item.style, "name", "body")
    text = to_plain_text(item.text)
    if not text:
        return

    if style_name in HEADING_STYLES:
        document.add_paragraph(text, style=HEADING_STYLES[style_name])
    elif style_name == "subtitle":
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif style_name == "bullet":
        document.add_paragraph(text, style="List Bullet")
    elif style_name == "caption":
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.italic = True
            run.font.size = Pt(9)
    else:
        document.add_paragraph(text)


def add_table(document: Document, item: Table) -> None:
    """Write a report table as a plain Word grid."""
    rows = getattr(item, "_cellvalues", None)
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            raw = cell.text if isinstance(cell, Paragraph) else str(cell)
            target = table.cell(r, c)
            target.text = to_plain_text(raw)
            for paragraph in target.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if r == 0:
                        run.bold = True


def add_code(document: Document, item: Preformatted) -> None:
    """Write a code listing as monospace lines, indentation preserved."""
    lines = getattr(item, "lines", None) or []
    for line in lines:
        text = line if isinstance(line, str) else " ".join(line)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text.replace("\t", "    "))
        run.font.name = "Courier New"
        run.font.size = Pt(8)


def add_image(document: Document, item: Image) -> None:
    """Embed a figure, scaled to the text width."""
    path = getattr(item, "filename", None)
    if not path or not Path(str(path)).exists():
        return
    width = min(MAX_IMAGE_WIDTH_IN, item.drawWidth / 72.0)
    document.add_picture(str(path), width=Inches(width))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> int:
    """Render the report story into a plain .docx."""
    story = flatten(build_story(load_artifacts()))
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    counts = {"paragraphs": 0, "tables": 0, "listings": 0, "figures": 0}
    for item in story:
        if isinstance(item, (Spacer, PageBreak)):
            continue
        if isinstance(item, Preformatted):
            add_code(document, item)
            counts["listings"] += 1
        elif isinstance(item, Paragraph):
            add_paragraph(document, item)
            counts["paragraphs"] += 1
        elif isinstance(item, Table):
            add_table(document, item)
            counts["tables"] += 1
        elif isinstance(item, Image):
            add_image(document, item)
            counts["figures"] += 1

    document.save(OUT)
    print(f"Wrote {OUT}")
    print("  " + ", ".join(f"{v} {k}" for k, v in counts.items()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
